def load_workbook(self, filepath_or_buffer: FilePathOrBuffer):
    from openpyxl import load_workbook
    return load_workbook(filepath_or_buffer, read_only=True, data_only=True, keep_links=False)